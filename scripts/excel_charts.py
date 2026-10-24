import datetime as dt
import posixpath
import re
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
}
COLORS = ["#3dd6b6", "#6e9fff", "#f4b942", "#ff758f", "#a989ff", "#90a5b3", "#70ad47"]


def empty_chart():
    return {"labels": [], "datasets": []}


def number(value):
    if isinstance(value, bool):
        return None
    try:
        result = float(value)
        return result if result == result and abs(result) != float("inf") else None
    except (TypeError, ValueError):
        return None


def label(value, monthly=False):
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime("%Y-%m" if monthly else "%m-%d")
    text = str(value or "")
    if monthly:
        match = re.search(r"(\d{1,2})", text)
        return f"{int(match.group(1)):02d}" if match else text
    match = re.match(r"^(\d{1,2})[-/月](\d{1,2})", text)
    return f"{int(match.group(1)):02d}-{int(match.group(2)):02d}" if match else text


def series_name(value):
    match = re.search(r"20\d{2}", str(value or ""))
    return match.group(0) + "年" if match else str(value or "系列")


def relationships(archive, path):
    rel_path = posixpath.join(posixpath.dirname(path), "_rels", posixpath.basename(path) + ".rels")
    if rel_path not in archive.namelist():
        return {}
    root = ET.fromstring(archive.read(rel_path))
    return {
        item.attrib["Id"]: posixpath.normpath(posixpath.join(posixpath.dirname(path), item.attrib["Target"]))
        for item in root
    }


def chart_series(path, sheet_name, position):
    with zipfile.ZipFile(path) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        workbook_rels = relationships(archive, "xl/workbook.xml")
        sheets = {
            item.attrib["name"]: workbook_rels[item.attrib["{" + NS["r"] + "}id"]]
            for item in workbook.find("m:sheets", NS)
        }
        sheet_path = sheets[sheet_name]
        sheet = ET.fromstring(archive.read(sheet_path))
        drawing_ref = sheet.find("m:drawing", NS)
        if drawing_ref is None:
            raise RuntimeError(f"工作表无图形 {sheet_name}")
        drawing_path = relationships(archive, sheet_path)[drawing_ref.attrib["{" + NS["r"] + "}id"]]
        drawing = ET.fromstring(archive.read(drawing_path))
        drawing_rels = relationships(archive, drawing_path)
        for anchor in drawing:
            start = anchor.find("xdr:from", NS)
            if start is None:
                continue
            col = int(start.find("xdr:col", NS).text) + 1
            row = int(start.find("xdr:row", NS).text) + 1
            if get_column_letter(col) + str(row) != position:
                continue
            chart_ref = anchor.find(".//c:chart", NS)
            chart_path = drawing_rels[chart_ref.attrib["{" + NS["r"] + "}id"]]
            chart = ET.fromstring(archive.read(chart_path))
            result = []
            for series in chart.findall(".//c:ser", NS):
                def formula(xpath):
                    node = series.find(xpath, NS)
                    return node.text if node is not None else None

                result.append({
                    "name": formula("c:tx/c:strRef/c:f"),
                    "cat": formula("c:cat/c:numRef/c:f") or formula("c:cat/c:strRef/c:f"),
                    "val": formula("c:val/c:numRef/c:f"),
                })
            if result:
                return result
        raise RuntimeError(f"原生图未找到 {sheet_name}!{position}")


def split_formula(formula):
    if not formula:
        return None
    match = re.match(r"^(?:'((?:[^']|'')+)'|([^!]+))!(.+)$", formula)
    if not match:
        return None
    sheet = (match.group(1) or match.group(2)).replace("''", "'")
    return sheet, match.group(3).replace("$", "")


class ChartExtractor:
    def __init__(self, path):
        self.path = path
        self.workbook = load_workbook(path, read_only=True, data_only=True)
        self.cache = {}

    def close(self):
        self.workbook.close()

    def cells(self, formula):
        if formula in self.cache:
            return self.cache[formula]
        parsed = split_formula(formula)
        if not parsed:
            return []
        sheet, ref = parsed
        min_col, min_row, max_col, max_row = range_boundaries(ref)
        values = []
        for row in self.workbook[sheet].iter_rows(
            min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col, values_only=True
        ):
            values.extend(row)
        self.cache[formula] = values
        return values

    def extract(self, sheet_name, position, config):
        definitions = chart_series(self.path, sheet_name, position)
        keep = set(config.get("keep_years") or [])
        scale = config.get("value_scale", 1)
        seasonal = config.get("seasonal", True)
        x_min = config.get("x_min")
        labels_out = []
        datasets = []
        records = []
        if config.get("add_series"):
            add = config["add_series"]
            definitions.append({
                "name": f"'{add.get('data_sheet', sheet_name)}'!$A$1",
                "cat": f"'{add.get('data_sheet', sheet_name)}'!{add['cat']}",
                "val": f"'{add.get('data_sheet', sheet_name)}'!{add['val']}",
                "literal_name": add["name"],
            })
        for index, definition in enumerate(definitions):
            name = definition.get("literal_name") or series_name(self.cells(definition["name"])[0] if definition["name"] else "")
            year_match = re.search(r"20\d{2}", name)
            if keep and year_match and int(year_match.group()) not in keep:
                continue
            values = self.cells(definition["val"])
            categories = self.cells(definition["cat"]) if definition["cat"] else []
            if config.get("x_labels") == "week":
                categories = [f"第{i + 1}周" for i in range(len(values))]
            elif not categories:
                categories = [f"{i + 1}月" for i in range(len(values))]
            points = []
            for i, value in enumerate(values):
                numeric = number(value)
                if numeric is None:
                    continue
                if seasonal and numeric == 0:
                    continue
                cat = categories[i] if i < len(categories) else ""
                if isinstance(cat, (dt.date, dt.datetime)) and x_min and cat < dt.datetime.fromisoformat(x_min):
                    continue
                text = label(cat, monthly="月" in str(cat)) if seasonal else (
                    cat.strftime("%Y-%m-%d") if isinstance(cat, (dt.date, dt.datetime)) else str(cat)
                )
                points.append((text, numeric * scale, cat))
            if not points:
                continue
            for point in points:
                if point[0] not in labels_out:
                    labels_out.append(point[0])
            datasets.append({
                "label": name,
                "data": [value for _, value, _ in points],
                "borderColor": COLORS[len(datasets) % len(COLORS)],
                "borderWidth": 2.2 if year_match and int(year_match.group()) == 2026 else 1.4,
            })
            records.append({"name": name, "points": points})
        for dataset, record in zip(datasets, records):
            lookup = {text: value for text, value, _ in record["points"]}
            dataset["data"] = [lookup.get(text) for text in labels_out]
        return {"labels": labels_out, "datasets": datasets}, records
